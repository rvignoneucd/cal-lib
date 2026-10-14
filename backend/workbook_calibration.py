"""Bounded workbook import; never execute formulas or infer compound identity."""
import base64
import hashlib
import io
import math
import re
import uuid
import zipfile

from openpyxl import load_workbook
from .gcms_calibration import fit_points

INSTRUMENTS = ('GC-MS', 'UV-VIS', 'LC-MS', 'Hapsite')
UNITS = ('ug/mL', 'mg/mL', 'ng/mL', 'mg/L', 'ug/L', 'mol/L', 'mmol/L')


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def fit_curve(curve):
    if not isinstance(curve, dict):
        raise ValueError('Invalid curve')
    name = curve.get('compound', '')
    if not isinstance(name, str) or not name.strip() or len(name) > 240:
        raise ValueError('Enter a compound name')
    if curve.get('unit') not in UNITS:
        raise ValueError('Select a concentration unit')
    points = curve.get('points')
    if not isinstance(points, list) or not 1 <= len(points) <= 200:
        raise ValueError('Each curve requires 1 to 200 rows')
    for point in points:
        if not isinstance(point, dict) or any(not number(point.get(k)) or point[k] < 0 for k in ('concentration', 'area')):
            raise ValueError('Concentrations and responses must be finite nonnegative numbers')
        if type(point.get('excluded', False)) is not bool:
            raise ValueError('Invalid point selection')
    rt = curve.get('rt')
    if rt is not None and (not number(rt) or rt < 0):
        raise ValueError('Invalid retention time')
    if type(curve.get('force_zero', False)) is not bool:
        raise ValueError('Invalid intercept setting')
    result = {**curve, 'compound': name.strip(), 'fit': None, 'issue': ''}
    try:
        result['fit'] = fit_points(points, curve.get('force_zero', False), auto_zero=not curve.get('force_zero_manual', False))
        result['force_zero'] = result['fit']['force_zero']
    except ValueError as error:
        result['issue'] = str(error)
    return result


def parse_workbook(filename, raw):
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('Use an .xlsx workbook; save legacy .xls files as .xlsx first')
    if len(raw) > 5_000_000:
        raise ValueError('Workbook exceeds 5 MB')
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 50_000_000:
            raise ValueError('Expanded workbook exceeds 50 MB')
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    curves, warnings = [], []
    digest = hashlib.sha256(raw).hexdigest()
    try:
        if len(book.worksheets) > 50:
            raise ValueError('Workbook exceeds 50 sheets')
        for sheet in book:
            if sheet.max_row > 10000 or sheet.max_column > 100:
                raise ValueError(f'{sheet.title}: sheet exceeds 10,000 rows or 100 columns')
            rows = list(sheet.values)
            for index, row in enumerate(rows):
                for col, value in enumerate(row):
                    if not isinstance(value, str) or not value.strip().lower().startswith('concentration'):
                        continue
                    response_col = next((j for j in range(col + 1, min(len(row), col + 4))
                                         if isinstance(row[j], str) and re.search(r'peak\s*area|absorbance|response|height', row[j], re.I)), None)
                    if response_col is None:
                        continue
                    context = rows[max(0, index - 5):index]
                    name = next((r[0].strip() for r in reversed(context) if isinstance(r[0], str)
                                 and r[0].strip() and not r[0].strip().lower().startswith('file')), '')
                    source = next((r[0].strip() for r in reversed(context) if isinstance(r[0], str)
                                   and r[0].strip().lower().startswith('file')), '')
                    rt = next((r[col] for r in reversed(context) if len(r) > col and number(r[col])), None)
                    unit = next((u for u in UNITS if u.lower() in value.replace('µ', 'u').replace('μ', 'u').lower()), '')
                    points, notes = [], []
                    for offset, values in enumerate(rows[index + 1:], index + 2):
                        x, y = values[col], values[response_col]
                        if isinstance(x, str) or (isinstance(values[0], str) and values[0].strip()):
                            break
                        if x is None and y is None:
                            continue
                        if not number(x) or not number(y) or x < 0 or y < 0:
                            warnings.append(f'{sheet.title} row {offset}: incomplete or invalid pair omitted')
                            continue
                        note = '; '.join(str(v) for v in values[response_col + 1:] if v is not None)
                        points.append({'concentration': x, 'area': y, 'excluded': False, 'source_row': offset, 'note': note})
                    for values in context:
                        notes.extend(str(v) for v in values[response_col + 1:] if v is not None)
                    if not points:
                        warnings.append(f'{sheet.title} row {index + 1}: no numeric pairs')
                        continue
                    curve = {'id': str(uuid.uuid5(uuid.NAMESPACE_URL, f'{digest}:{sheet.title}:{index}:{col}')),
                             'compound': name or sheet.title, 'rt': rt, 'unit': unit or 'ug/mL',
                             'source': source, 'workbook': filename, 'sheet': sheet.title, 'header_row': index + 1,
                             'notes': '; '.join(notes), 'response_type': str(row[response_col]).strip(),
                             'points': points, 'force_zero': False, 'selected': True, 'reviewed': False}
                    if not unit:
                        curve['notes'] += ' Concentration unit not recognized; verify the selected unit.'
                    curves.append(fit_curve(curve))
                    if len(curves) > 100:
                        raise ValueError('Workbook exceeds 100 curves')
    finally:
        book.close()
    if not curves:
        raise ValueError('No concentration / response blocks found')
    return curves, warnings


def preview_upload(payload):
    instrument = payload.get('instrument')
    files = payload.get('files')
    if instrument not in INSTRUMENTS or not isinstance(files, list) or not 1 <= len(files) <= 10:
        raise ValueError('Select an instrument and 1 to 10 workbooks')
    curves, warnings, errors, seen = [], [], [], set()
    if sum(len(f.get('data', '')) for f in files) > 28_000_000:
        raise ValueError('Combined upload exceeds 20 MB')
    for file in files:
        name = str(file.get('filename', ''))[:240]
        try:
            raw = base64.b64decode(file.get('data', ''), validate=True)
            imported, issues = parse_workbook(name, raw)
            for curve in imported:
                if curve['id'] not in seen:
                    seen.add(curve['id']); curves.append(curve)
            warnings.extend(f'{name}: {issue}' for issue in issues)
        except Exception as error:
            errors.append(f'{name}: {error}')
    if len(curves) > 100:
        raise ValueError('Import up to 100 curves at a time')
    return {'id': str(uuid.uuid4()), 'instrument': instrument, 'curves': curves,
            'warnings': warnings, 'errors': errors, 'revision': 0, 'status': 'draft'}
