"""Compact calibration table with an editable, cell-linked Excel scatter chart."""
import math

from fastapi import HTTPException
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.data_source import NumData, NumVal, NumFmt
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.trendline import Trendline, TrendlineLabel
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties


def _text(value):
    text = str(value or '')
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) else text


def format_calibration_export(workbook, source, payload):
    rows = payload.get('rows') or []
    calibration = payload.get('calibration') or {}
    if not isinstance(calibration, dict) or not 3 <= len(rows) <= 1000:
        raise HTTPException(400, 'Calibration export requires at least three measured standards')
    for row in rows:
        if not isinstance(row, dict) or any(
            isinstance(row.get(key), bool) or not isinstance(row.get(key), (int, float))
            or not math.isfinite(row[key]) or row[key] < 0 for key in ('concentration', 'area')
        ):
            raise HTTPException(400, 'Calibration export requires numeric concentrations and areas')
    if len({row['concentration'] for row in rows}) < 3:
        raise HTTPException(400, 'Calibration export requires three different concentrations')
    if len({row['area'] for row in rows}) < 2:
        raise HTTPException(400, 'Calibration areas must vary')
    ordered = sorted(rows, key=lambda row: row['concentration'])
    name = _text(calibration.get('name') or rows[0].get('compound') or 'Compound')
    unit = _text(calibration.get('unit') or rows[0].get('unit') or '')
    force_zero = calibration.get('force_zero') is True
    source.title = 'Source data'
    # Generic peak exports preserve area as text. Calibration charts require numbers.
    for column_index, column in enumerate(payload.get('columns', []), 1):
        if column.get('key') == 'area':
            for row_index, row in enumerate(rows, 2):
                source.cell(row_index, column_index, row['area']).number_format = '0.00E+00'
    source.auto_filter.ref = source.dimensions
    sheet = workbook.create_sheet('Calibration', 0)
    workbook.active = 0
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'B2'
    sheet.append([name, 'conc', 'peak area', 'Notes'])
    for row in ordered:
        sheet.append([None, row['concentration'], row['area'], _text(row.get('notes'))])
    for column, width in {'A':22, 'B':14, 'C':21, 'D':24, 'E':3,
                           'F':24, 'G':23, 'H':16, 'I':16, 'J':16, 'K':8}.items():
        sheet.column_dimensions[column].width = width
    divider = Side(style='thin', color='D9DFE3')
    for row in sheet.iter_rows(min_row=1, max_row=len(rows)+1, max_col=4):
        for cell in row:
            cell.font = Font(name='Arial', size=12, color='18252D', bold=cell.row == 1)
            cell.alignment = Alignment(vertical='center', wrap_text=cell.column in (1,4))
            cell.border = Border(bottom=divider)
            if cell.row == 1:
                cell.fill = PatternFill('solid', fgColor='E5ECEF')
        sheet.row_dimensions[row[0].row].height = 25
    for row in range(2, len(rows)+2):
        sheet.cell(row,2).number_format = '0.########'
        sheet.cell(row,3).number_format = '0.00E+00'
    end = len(rows) + 1
    xs, ys = f'B2:B{end}', f'C2:C{end}'
    notes = [
        ('Notes', None),
        ('Concentration unit', unit),
        ('Slope', f'=SUMPRODUCT({xs},{ys})/SUMSQ({xs})' if force_zero else f'=SLOPE({ys},{xs})'),
        ('Intercept', '=0' if force_zero else f'=INTERCEPT({ys},{xs})'),
        ('R2 (centered)', f'=1-SUMPRODUCT(({ys}-($G$3*{xs}+$G$4))^2)/DEVSQ({ys})' if force_zero else f'=RSQ({ys},{xs})'),
        ('Fit', 'Through zero' if force_zero else 'Free intercept'),
    ]
    for index, (label, value) in enumerate(notes, 1):
        sheet.cell(index,6,label).font = Font(name='Arial', size=12, bold=index==1)
        sheet.cell(index,7,value).font = Font(name='Arial', size=12)
        sheet.row_dimensions[index].height = 25
    sheet['G3'].number_format = sheet['G4'].number_format = '0.00000E+00'
    sheet['G5'].number_format = '0.00000'
    chart = ScatterChart()
    chart.title = f'{name} calibration'
    chart.x_axis.title = f'Concentration ({unit})'
    chart.y_axis.title = 'Peak area'
    chart.x_axis.scaling.min = 0
    chart.y_axis.scaling.min = 0
    chart.y_axis.numFmt = '0.00E+00'
    chart.width, chart.height = 21, 13
    chart.legend = None
    chart.scatterStyle = 'marker'
    series = Series(Reference(sheet, min_col=3, min_row=2, max_row=end),
                    Reference(sheet, min_col=2, min_row=2, max_row=end), title='Standards')
    series.marker.symbol = 'circle'
    series.marker.size = 7
    series.marker.graphicalProperties.solidFill = '176784'
    series.marker.graphicalProperties.line.solidFill = '176784'
    series.graphicalProperties.line.noFill = True
    trend_style = GraphicalProperties()
    trend_style.line.solidFill = '176784'
    trend_style.line.prstDash = 'sysDot'
    # Excel's through-zero R2 convention can differ; the adjacent formula is centered.
    series.trendline = Trendline(trendlineType='linear', intercept=0 if force_zero else None,
                                dispEq=True, dispRSqr=not force_zero, spPr=trend_style,
                                trendlineLbl=TrendlineLabel(numFmt=NumFmt(formatCode='0.00E+00',sourceLinked=False),
                                    layout=Layout(manualLayout=ManualLayout(x=.08,y=.06))))
    for ref, key in ((series.xVal.numRef, 'concentration'), (series.yVal.numRef, 'area')):
        ref.numCache = NumData(ptCount=len(ordered), pt=[NumVal(idx=i,v=row[key]) for i,row in enumerate(ordered)])
    chart.series.append(series)
    sheet.add_chart(chart, 'F8')
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f'A1:N{max(35,end)}'
