"""Editable calibration workbooks for the private lab."""
from io import BytesIO
import math
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font
from .calibration_export import format_calibration_export, _text


def calibration_workbook(payload):
    curves = payload.get('curves', [])
    if not isinstance(curves, list) or not 1 <= len(curves) <= 200:
        raise HTTPException(400, 'Choose one to 200 compounds')
    book = Workbook()
    caches = {}
    book.remove(book.active)
    for index, curve in enumerate(curves, 1):
        rows = curve.get('rows', [])
        if not isinstance(rows, list) or not 1 <= len(rows) <= 1000:
            raise HTTPException(400, 'No measurements to export')
        for row in rows:
            for key in ('concentration', 'area', 'rt', 'wavelength'):
                value = row.get(key)
                if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0):
                    raise HTTPException(400, 'Invalid measurement')
        source = book.create_sheet(f'Data {index}')
        source.append(['Compound', 'Concentration', 'Response', 'RT (min)', 'Wavelength (nm)', 'Source', 'Unit', 'Status'])
        for row in rows:
            source.append([_text(curve.get('name')), row.get('concentration'), row.get('area'), row.get('rt'), row.get('wavelength'), _text(row.get('source')), _text(curve.get('unit')), _text(curve.get('status'))])
        source.freeze_panes = 'A2'
        for cell in source[1]:
            cell.font = Font(bold=True)
        for col in 'ABCDEFGH':
            source.column_dimensions[col].width = 24
        source.auto_filter.ref = source.dimensions
        measured = [r for r in rows if r.get('concentration') is not None and r.get('area') is not None]
        if len({r['concentration'] for r in measured}) >= 3 and len({r['area'] for r in measured}) >= 2:
            # Build each curve separately so fixed formula references stay local.
            temporary = Workbook()
            format_calibration_export(temporary, temporary.active, {'rows': measured, 'calibration': curve})
            original = temporary['Calibration']
            sheet = book.create_sheet(f'Curve {index}')
            cached = caches.setdefault(f'xl/worksheets/sheet{book.index(sheet)+1}.xml', {})
            from copy import copy
            for row in original:
                for cell in row:
                    target = sheet.cell(cell.row, cell.column, cell.value)
                    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
                        setattr(target, attribute, copy(getattr(cell, attribute)))
                    target.number_format = cell.number_format
            for key, dimension in original.column_dimensions.items():
                sheet.column_dimensions[key].width = dimension.width
            # Reformat the existing cell-linked exporter to the supplied GC worksheet.
            from openpyxl.styles import Alignment
            from openpyxl.chart import Reference
            from openpyxl.chart.data_source import NumData, NumVal
            sheet.delete_rows(1, sheet.max_row)
            sheet.sheet_view.showGridLines = True
            for col, width in {'A':24, 'B':23, 'C':15, 'D':12, 'E':10}.items():
                sheet.column_dimensions[col].width = width
            sheet['B1'] = 'RT (min)'
            sheet['A2'] = _text(curve.get('name'))
            sheet['A2'].font = Font(name='Calibri', size=11, bold=True)
            sheet['A2'].alignment = Alignment(wrap_text=True, vertical='top')
            sheet['B2'] = f'=IF(COUNT(\'Data {index}\'!D2:D{len(rows)+1})=0,"",AVERAGE(\'Data {index}\'!D2:D{len(rows)+1}))'
            sheet['B2'].number_format = '0.000'
            rts = [row['rt'] for row in rows if row.get('rt') is not None]
            if rts:
                cached['B2'] = sum(rts)/len(rts)
            notes = ' | '.join(str(v) for v in (curve.get('notes'), curve.get('status')) if v)
            sheet.merge_cells('D2:P2')
            sheet['D2'] = _text(notes)
            sheet['D2'].alignment = Alignment(wrap_text=True, vertical='top')
            sheet.row_dimensions[2].height = max(30, 15*math.ceil(len(notes)/120), 15*math.ceil(len(str(curve.get('name','')))/24))
            sheet['B4'] = _text(f"Concentration ({curve.get('unit') or ''})")
            sheet['C4'] = _text(curve.get('response_label') or 'Peak Area')
            for n, row in enumerate(sorted(measured, key=lambda r: r['concentration']), 5):
                source_row = rows.index(row) + 2
                for target, origin in [('B', 'B'), ('C', 'C')]:
                    sheet[f'{target}{n}'] = f"=IF('Data {index}'!{origin}{source_row}=\"\",\"\",'Data {index}'!{origin}{source_row})"
                    cached[f'{target}{n}'] = row['concentration' if target == 'B' else 'area']
                sheet[f'C{n}'].number_format = '0.00E+00'
            end = len(measured)+4
            xs, ys = f'B5:B{end}', f'C5:C{end}'
            slope_row = end+2
            zero = curve.get('force_zero') is True
            mx = sum(row['concentration'] for row in measured)/len(measured)
            my = sum(row['area'] for row in measured)/len(measured)
            slope = (sum(row['concentration']*row['area'] for row in measured)/sum(row['concentration']**2 for row in measured) if zero else
                     sum((row['concentration']-mx)*(row['area']-my) for row in measured)/sum((row['concentration']-mx)**2 for row in measured))
            intercept = 0 if zero else my-slope*mx
            r2 = 1-sum((row['area']-slope*row['concentration']-intercept)**2 for row in measured)/sum((row['area']-my)**2 for row in measured)
            cached.update({f'C{slope_row}': slope, f'C{slope_row+1}': intercept, f'C{slope_row+2}': r2})
            for offset, (label, formula) in enumerate([
                ('Slope', f'=SUMPRODUCT({xs},{ys})/SUMSQ({xs})' if zero else f'=SLOPE({ys},{xs})'),
                ('Intercept', '=0' if zero else f'=INTERCEPT({ys},{xs})'),
                ('R2 (centered)', f'=1-SUMPRODUCT(({ys}-(C{slope_row}*{xs}+C{slope_row+1}))^2)/DEVSQ({ys})' if zero else f'=RSQ({ys},{xs})'),
            ]):
                sheet.cell(slope_row+offset,2,label)
                sheet.cell(slope_row+offset,3,formula).number_format = '0.0000' if offset==2 else '0.00'
            chart = original._charts[0]
            for series in chart.series:
                for ref in (series.xVal.numRef, series.yVal.numRef):
                    column = 2 if ref is series.xVal.numRef else 3
                    ref.f = str(Reference(sheet,min_col=column,min_row=5,max_row=end))
                series.marker.size = 5
                series.marker.graphicalProperties.solidFill = '4472C4'
                series.marker.graphicalProperties.line.solidFill = '4472C4'
                series.trendline.spPr.line.solidFill = '4472C4'
            kind = 'IS-normalized' if curve.get('response_label')=='Area ratio' else 'External'
            chart.title = f"{curve.get('name') or 'Compound'} {kind} Cal"
            chart.x_axis.title = chart.y_axis.title = None
            chart.width, chart.height = 16, 9.5
            sheet.add_chart(chart, 'F4')
            sheet.freeze_panes = None
            book.calculation = temporary.calculation
            if index == 1:
                book.active = book.index(sheet)
    stream = BytesIO()
    book.save(stream)
    # Keep editable formulas while supplying results to non-recalculating viewers.
    output = BytesIO()
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    with ZipFile(stream) as source, ZipFile(output, 'w') as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename in caches:
                root = ET.fromstring(content)
                for cell in root.findall(f'.//{{{ns}}}c'):
                    value = caches[item.filename].get(cell.get('r'))
                    if value is not None:
                        cached_value = cell.find(f'{{{ns}}}v')
                        if cached_value is None:
                            cached_value = ET.SubElement(cell, f'{{{ns}}}v')
                        cached_value.text = repr(value)
                content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(item, content)
    return output.getvalue()
